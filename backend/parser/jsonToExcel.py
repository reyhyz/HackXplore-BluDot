"""
jsonToExcel.py  (updated)
-------------------------
Two additions vs the original:
  1. output_path now accepts a str, Path, OR a file-like object (e.g. io.BytesIO)
     so the FastAPI server can stream the file without touching disk.
  2. Data blocks with "_section_header": True get a grey divider row instead of
     the normal green/red colouring, making multi-batch exports readable.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Union, IO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


_GREEN_FILL    = PatternFill("solid", start_color="C6EFCE")
_RED_FILL      = PatternFill("solid", start_color="FFC7CE")
_HEADER_FILL   = PatternFill("solid", start_color="D9E1F2")
_META_FILL     = PatternFill("solid", start_color="F2F2F2")
_SECTION_FILL  = PatternFill("solid", start_color="D6DCE4")   # blue-grey divider

_HEADER_FONT  = Font(name="Arial", bold=True,  size=10, color="1F3864")
_BODY_FONT    = Font(name="Arial",             size=10)
_META_FONT    = Font(name="Arial", italic=True, size=10, color="595959")
_TITLE_FONT   = Font(name="Arial", bold=True,  size=13, color="1F3864")
_SECTION_FONT = Font(name="Arial", bold=True,  size=10, color="374151")

_THIN_BORDER = Border(
    left=Side(style="thin",   color="BFBFBF"),
    right=Side(style="thin",  color="BFBFBF"),
    top=Side(style="thin",    color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)

_COLUMNS: list[tuple[str, str, int]] = [
    ("Parameter",      "parameter",      32),
    ("Expected Value", "expected_value", 24),
    ("Actual Value",   "actual_value",   24),
    ("Page",           "page_number",     8),
    ("Line",           "line_number",     8),
    ("Flags",          "flags",          40),
    ("Status",         "status",         14),
]


def _flags_to_str(flags: dict) -> str:
    return "\n".join(f"{k}: {v}" for k, v in flags.items()) if flags else ""


def _write_meta_block(ws, meta: dict[str, str]) -> int:
    ws.merge_cells("A1:G1")
    c = ws["A1"]
    c.value = "Batch Review Report"
    c.font  = _TITLE_FONT
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    row = 2
    for key, value in meta.items():
        for col, val in ((1, key), (2, value)):
            cell = ws.cell(row=row, column=col, value=val)
            cell.fill = _META_FILL
            cell.font = _META_FONT
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 16
        row += 1

    row += 1   # blank spacer
    return row


def _write_header_row(ws, row: int) -> None:
    for col_idx, (label, _, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=row, column=col_idx, value=label)
        cell.font      = _HEADER_FONT
        cell.fill      = _HEADER_FILL
        cell.border    = _THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[row].height = 20


def _write_section_header_row(ws, row: int, label: str) -> None:
    """Grey divider row used between batches in a multi-batch export."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(_COLUMNS))
    cell = ws.cell(row=row, column=1, value=label)
    cell.fill      = _SECTION_FILL
    cell.font      = _SECTION_FONT
    cell.border    = _THIN_BORDER
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 18


def _write_data_row(ws, row: int, block: dict[str, Any]) -> None:
    flags       = block.get("flags") or {}
    row_fill    = _GREEN_FILL if not flags else _RED_FILL
    status_text = "OK" if not flags else f"{len(flags)} issue(s)"

    values = [
        block.get("parameter",      ""),
        block.get("expected_value", ""),
        block.get("actual_value",   ""),
        block.get("page_number",    ""),
        block.get("line_number",    ""),
        _flags_to_str(flags),
        status_text,
    ]

    row_height = max(15 * len(flags), 15) if flags else 15

    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col_idx, value=value)
        cell.fill      = row_fill
        cell.font      = _BODY_FONT
        cell.border    = _THIN_BORDER
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
            horizontal="center" if col_idx in (4, 5) else "left",
        )

    ws.row_dimensions[row].height = row_height


def export_to_excel(
    data_blocks:  list[dict[str, Any]],
    output_path:  Union[str, Path, IO[bytes]] = "batch_review.xlsx",
    sheet_name:   str = "Batch Review",
    batch_meta:   dict[str, str] | None = None,
) -> Union[Path, None]:
    """
    Export data_blocks to Excel.

    output_path can be:
      - a str / Path  →  writes to disk, returns the Path
      - a file-like   →  writes into the buffer (e.g. io.BytesIO), returns None
    """
    is_filelike = hasattr(output_path, "write")

    if not is_filelike:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.sheet_view.showGridLines = False

    next_row = 1
    if batch_meta:
        next_row = _write_meta_block(ws, batch_meta)

    _write_header_row(ws, next_row)
    freeze_at = next_row + 1       # freeze everything above the first data row
    next_row += 1

    real_data_blocks = []          # exclude section headers from summary count

    for block in data_blocks:
        if block.get("_section_header"):
            _write_section_header_row(ws, next_row, block.get("parameter", ""))
        else:
            _write_data_row(ws, next_row, block)
            real_data_blocks.append(block)
        next_row += 1

    # Summary row
    total      = len(real_data_blocks)
    ok_count   = sum(1 for b in real_data_blocks if not (b.get("flags") or {}))
    flag_count = total - ok_count

    ws.cell(row=next_row, column=1, value="SUMMARY").font = _HEADER_FONT
    ws.cell(row=next_row, column=2, value=f"Total: {total}").font   = _BODY_FONT
    ws.cell(row=next_row, column=3, value=f"OK: {ok_count}").font   = Font(name="Arial", size=10, color="375623")
    ws.cell(row=next_row, column=4, value=f"Issues: {flag_count}").font = Font(name="Arial", size=10, color="9C0006")

    for col in range(1, len(_COLUMNS) + 1):
        cell = ws.cell(row=next_row, column=col)
        cell.fill   = _META_FILL
        cell.border = _THIN_BORDER

    ws.freeze_panes = ws.cell(row=freeze_at, column=1)

    wb.save(output_path)
    return None if is_filelike else output_path


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("Usage: python jsonToExcel.py input.json [output.xlsx]")
        sys.exit(1)

    input_path  = Path(sys.argv[1])
    output_path = Path(sys.argv[2]) if len(sys.argv) > 2 else input_path.with_suffix(".xlsx")

    raw = json.loads(input_path.read_text(encoding="utf-8"))

    if isinstance(raw, list):
        blocks, meta = raw, None
    else:
        blocks = raw.get("extracted_data", [])
        meta   = raw.get("meta")

    out = export_to_excel(blocks, output_path=output_path, batch_meta=meta)
    print(f"Exported → {out}")